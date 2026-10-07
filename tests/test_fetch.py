"""Get the data, no code: `ghost-api fetch` and the dashboard's Download data (phase 3)."""

import csv
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from typer.testing import CliRunner

from ghost_api.cli import app
from ghost_api.export.tabular import columns, to_bytes
from ghost_api.memory.store import Store
from ghost_api.web.server import RunManager, create_app
from tests.test_runtime import shop_db  # noqa: F401  (module-scoped demo-shop run)

runner = CliRunner()


def test_fetch_lists_what_can_be_fetched(shop_db: tuple[Path, str]) -> None:  # noqa: F811
    db, _ = shop_db
    result = runner.invoke(app, ["fetch", "127.0.0.1", "--db", str(db)])
    assert result.exit_code == 0, result.output
    assert "get_products" in result.output and "per_page=" in result.output
    assert "ghost-api fetch 127.0.0.1" in result.output


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_fetch_downloads_all_pages(
    shop_db: tuple[Path, str],  # noqa: F811
    tmp_path: Path,
    fmt: str,
) -> None:
    db, _ = shop_db
    target = tmp_path / f"products.{fmt}"
    result = runner.invoke(
        app,
        ["fetch", "127.0.0.1", "get_products", "per_page=10", "--pages", "10",
         "--format", fmt, "--save-to", str(target), "--db", str(db)],
    )  # fmt: skip
    assert result.exit_code == 0, result.output
    assert "Saved 30 rows" in result.output and "Page 3: 10 rows" in result.output
    if fmt == "csv":
        rows = list(csv.DictReader(io.StringIO(target.read_text(encoding="utf-8-sig"))))
        assert len(rows) == 30 and {"id", "name", "category", "price"} <= set(rows[0])
    elif fmt == "xlsx":
        sheet = load_workbook(target).active
        assert sheet is not None and sheet.max_row == 31  # header + 30
    else:
        data = json.loads(target.read_text(encoding="utf-8"))
        assert len(data) == 30 and isinstance(data[0]["id"], int)


def test_fetch_explains_bad_input(shop_db: tuple[Path, str]) -> None:  # noqa: F811
    db, _ = shop_db
    unknown = runner.invoke(app, ["fetch", "127.0.0.1", "nope", "--db", str(db)])
    assert unknown.exit_code == 1 and "Choose one of" in unknown.output
    bad = runner.invoke(app, ["fetch", "127.0.0.1", "get_products", "page", "--db", str(db)])
    assert bad.exit_code == 2 and "name=value" in bad.output
    fmt = runner.invoke(app, ["fetch", "127.0.0.1", "get_products", "--format", "pdf",
                              "--db", str(db)])  # fmt: skip
    assert fmt.exit_code == 2 and "csv, xlsx, json" in fmt.output


def test_spreadsheets_never_get_formulas_and_nested_fields_flatten() -> None:
    records = [{"name": "=HYPERLINK(\"http://evil\")", "venue": {"area": "Velachery"},
                "tags": ["a", "b"], "price": -5}]  # fmt: skip
    text = to_bytes(records, "csv").decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "venue.area" in text and "a; b" in text and "-5" in text
    sheet = load_workbook(io.BytesIO(to_bytes(records, "xlsx"))).active
    assert sheet is not None and sheet["A2"].value.startswith("'=")
    assert columns([{"a": 1}, {"b": 2, "a": 3}]) == ["a", "b"]


def test_dashboard_download_and_allow(shop_db: tuple[Path, str], tmp_path: Path) -> None:  # noqa: F811
    db, _ = shop_db
    out = tmp_path / "out"
    client = TestClient(create_app(db, out, "t", manager=RunManager(db, out, lambda c: None)))
    auth = {"x-ghost-token": "t"}
    detail = client.get("/api/sites/127.0.0.1", headers=auth).json()
    products = next(e for e in detail["endpoints"] if e["label"] == "GET /api/products")
    assert products["callable"] is True and {"page", "per_page"} <= {
        i["name"] for i in products["inputs"]
    }
    resp = client.post(
        "/api/sites/127.0.0.1/fetch",
        headers=auth,
        json={"endpoint": products["key"], "params": {"per_page": "10"}, "pages": 10,
              "format": "csv"},
    )  # fmt: skip
    assert resp.status_code == 200, resp.text
    assert resp.headers["x-ghost-rows"] == "30"
    assert 'filename="127.0.0.1-get_products.csv"' in resp.headers["content-disposition"]
    bad = client.post("/api/sites/127.0.0.1/fetch", headers=auth,
                      json={"endpoint": products["key"], "format": "exe"})  # fmt: skip
    assert bad.status_code == 400
    # The allow switch stores the choice (and can take it back).
    ok = client.post("/api/sites/127.0.0.1/allow", headers=auth,
                     json={"endpoint": products["key"], "allowed": True})  # fmt: skip
    assert ok.status_code == 200
    with Store(db) as store:
        assert products["key"] in store.allowed_endpoints("127.0.0.1")
    client.post("/api/sites/127.0.0.1/allow", headers=auth,
                json={"endpoint": products["key"], "allowed": False})  # fmt: skip
    with Store(db) as store:
        assert products["key"] not in store.allowed_endpoints("127.0.0.1")
    missing = client.post("/api/sites/127.0.0.1/allow", headers=auth,
                          json={"endpoint": "nope", "allowed": True})  # fmt: skip
    assert missing.status_code == 404
